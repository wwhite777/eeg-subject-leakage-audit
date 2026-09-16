#!/usr/bin/env python3
"""Shared plotting helpers for the revision figures (exact data from results/revision CSVs only).

Design rules (dataviz + research-figures skills, EIC request): every panel shows the individual
subjects, labelled axes with units, a legend for >= 2 series, colourblind-safe fixed hue order
(validated with the dataviz palette validator on 2026-09-16), fonts designed for the printed width
(IOP text width 153 mm), and a caption written from the same CSVs. Markers never encode significance.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import common as C  # noqa: E402

FIG_DIR = C.ROOT / "result" / "figure" / "revision"

# fixed hue order (Okabe–Ito), validated: protocols and models never share a slot within one figure
PROTO_COLOR = {"within_subject": "#009E73", "within_session": "#009E73", "trial_random": "#D55E00", "trial_random_6": "#D55E00",
               "trial_random_2": "#CC79A7", "run_disjoint": "#E69F00", "session_disjoint": "#56B4E9", "subject_disjoint": "#0072B2",
               "loso": "#0072B2", "subject_disjoint_5": "#999999"}
PROTO_LABEL = {"within_subject": "within-subject", "within_session": "within-subject (session→session)", "trial_random": "trial-random (P0)",
               "trial_random_6": "trial-random 6-fold (P0a)", "trial_random_2": "trial-random 2-fold (P0b)", "run_disjoint": "run-disjoint (P1)",
               "session_disjoint": "session-disjoint (P1s)", "subject_disjoint": "subject-disjoint (P2)", "loso": "subject-disjoint LOSO (P2)",
               "subject_disjoint_5": "subject-disjoint 5-fold (original)"}
MODEL_COLOR = {"csp_lda": "#CC79A7", "riemann_ts_lr": "#0072B2", "shallow_fbcsp": "#E69F00", "eegnet_v4": "#009E73"}
MODEL_LABEL = {"csp_lda": "CSP+LDA", "riemann_ts_lr": "TS+LR", "shallow_fbcsp": "ShallowFBCSPNet", "eegnet_v4": "EEGNet"}
MODEL_MARK = {"csp_lda": "s", "riemann_ts_lr": "o", "shallow_fbcsp": "^", "eegnet_v4": "D"}


def setup_mpl():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({
        "font.size": 8, "axes.titlesize": 8.5, "axes.labelsize": 8, "legend.fontsize": 7, "xtick.labelsize": 7, "ytick.labelsize": 7,
        "font.family": "DejaVu Sans", "axes.spines.top": False, "axes.spines.right": False, "axes.grid": True, "grid.color": "#e6e6e6",
        "grid.linewidth": 0.5, "axes.axisbelow": True, "lines.linewidth": 1.4, "savefig.dpi": 300, "figure.dpi": 100,
    })
    return plt


def seed_avg_subject(rows, keys, subject_key="subject", acc_key="accuracy"):
    """{key_tuple: {subject: mean over seeds/draws}}"""
    tmp = defaultdict(lambda: defaultdict(list))
    for r in rows:
        k = tuple(r[x] for x in keys)
        tmp[k][int(r[subject_key])].append(float(r[acc_key]))
    return {k: {s: float(np.mean(v)) for s, v in d.items()} for k, d in tmp.items()}


def paradigm_mean(subj_by_key, key_index_task, tasks=C.TASKS):
    """Average per subject over the four EEGBCI paradigms (key element key_index_task is the task)."""
    agg = defaultdict(lambda: defaultdict(list))
    for k, d in subj_by_key.items():
        if k[key_index_task] not in tasks:
            continue
        nk = tuple("paradigm_mean" if i == key_index_task else x for i, x in enumerate(k))
        for s, v in d.items():
            agg[nk][s].append(v)
    return {k: {s: float(np.mean(v)) for s, v in d.items() if len(v) == len(tasks)} for k, d in agg.items()}


def strip_points(ax, x, values, color, rng, width=0.16, size=7, alpha=0.45, zorder=2, marker="o"):
    values = np.asarray(values, dtype=float)
    jitter = rng.uniform(-width, width, size=len(values))
    ax.scatter(x + jitter, values, s=size, color=color, alpha=alpha, linewidths=0, zorder=zorder, marker=marker, rasterized=True)


def mean_ci(ax, x, values, color, seed=C.SEED, width=0.28, lw=1.6, zorder=4, marker="_"):
    values = np.asarray(values, dtype=float)
    m = values.mean()
    lo, hi = C.bootstrap_ci(values, seed=seed)
    ax.plot([x - width, x + width], [m, m], color="black", lw=lw, zorder=zorder + 1, solid_capstyle="butt")
    ax.plot([x, x], [lo, hi], color="black", lw=1.0, zorder=zorder)
    return m, lo, hi


def save(fig, name: str, caption: str, data_files: list[Path]):
    """PNG (300 dpi) + PDF + caption .txt naming the data files; also a large-font pptx with the PNG."""
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    png = FIG_DIR / f"{name}.png"
    fig.savefig(png, bbox_inches="tight")
    fig.savefig(FIG_DIR / f"{name}.pdf", bbox_inches="tight")
    def rel(p: Path) -> str:
        p = Path(p).resolve()
        return str(p.relative_to(C.ROOT)) if str(p).startswith(str(C.ROOT)) else str(p)

    (FIG_DIR / f"{name}_caption.txt").write_text(caption + "\n\nData: " + ", ".join(rel(p) for p in data_files) + "\n", encoding="utf-8")
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.7))
        tb.text_frame.text = name.replace("_", " ")
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
        slide.shapes.add_picture(str(png), Inches(0.6), Inches(0.9), height=Inches(6.3))
        prs.save(str(FIG_DIR / f"{name}.pptx"))
    except Exception as exc:  # noqa: BLE001
        print("pptx skipped:", exc)
    print("saved", png)
    return png
